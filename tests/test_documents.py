import io
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from backend.documents import extract_document


def pdf_bytes(text="A readable document passage.", encrypted=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if text:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 250 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("password")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def word_bytes():
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Running header not narrative"
    doc.add_paragraph("Chapter one", style="Heading 1")
    doc.add_paragraph("First paragraph.")
    doc.add_table(rows=1, cols=2).cell(0, 0).text = "Table entry"
    doc.add_paragraph("Final paragraph.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_docx_body_order_and_header_exclusion(tmp_path):
    source = tmp_path / "book.docx"
    source.write_bytes(word_bytes())
    text = "\n\n".join(extract_document(source))
    assert text.index("Chapter one") < text.index("First paragraph") < text.index("Table entry") < text.index("Final paragraph")
    assert "Running header" not in text


@pytest.mark.parametrize("data,message", [(pdf_bytes(""), "No embedded text"), (pdf_bytes(encrypted=True), "Password-protected")])
def test_unusable_pdf_is_rejected(tmp_path, data, message):
    source = tmp_path / "book.pdf"
    source.write_bytes(data)
    with pytest.raises(ValueError, match=message):
        extract_document(source)


def test_multi_page_pdf_extracts_without_ocr(tmp_path):
    from pypdf import PdfReader
    source = tmp_path / "synthetic-book.pdf"
    writer = PdfWriter()
    for i in range(6):
        writer.add_page(PdfReader(io.BytesIO(pdf_bytes(f"Chapter {i + 1}. Original synthetic sample."))).pages[0])
    with source.open("wb") as output:
        writer.write(output)
    pages = extract_document(source)
    assert len(pages) == 6 and all(pages)
    assert "Chapter 1" in pages[0] and "Chapter 6" in pages[-1]
